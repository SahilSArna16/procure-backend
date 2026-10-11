import io
import zipfile
from pathlib import Path

import pdfplumber
from docx import Document
from openpyxl import load_workbook
from PIL import Image

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_TEXT = 100_000
SUPPORTED = {".xlsx", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".txt"}


def check_archive(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
            raise ValueError("Office file expands beyond the 50 MB limit.")


def bounded_sheet_rows(sheet):
    """Read actual cells: XLSX dimensions are optional and may also be stale."""
    sheet.reset_dimensions()
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index > 5000 or len(row) > 100:
            raise ValueError("Excel sheet exceeds 5,000 rows or 100 columns.")
        yield index, row


def parse_document(filename: str, data: bytes) -> str | None:
    """None requests OCR. Preserve tables, page markers and footnotes as text."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("Supported formats: XLSX, PDF, DOCX, PNG, JPG, TXT.")
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("Each file must be nonempty and at most 15 MB.")
    if suffix in {".png", ".jpg", ".jpeg"}:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
        return None
    if suffix == ".txt":
        text = data.decode("utf-8-sig")
    elif suffix == ".xlsx":
        check_archive(data)
        book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sections = []
        try:
            for sheet in book:
                sections.append(f"Sheet: {sheet.title}")
                for index, row in bounded_sheet_rows(sheet):
                    if any(value is not None for value in row):
                        sections.append(f"Row {index}: " + " | ".join(str(v) if v is not None else "" for v in row))
        finally:
            book.close()
        text = "\n".join(sections)
    elif suffix == ".docx":
        check_archive(data)
        doc = Document(io.BytesIO(data))
        sections = [p.text for p in doc.paragraphs if p.text]
        for index, table in enumerate(doc.tables, start=1):
            sections.append(f"Table {index}")
            sections.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        for section in doc.sections:
            sections.extend(p.text for p in section.header.paragraphs + section.footer.paragraphs if p.text)
        # Word footnotes/endnotes are not surfaced by python-docx.
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for name in ("word/footnotes.xml", "word/endnotes.xml"):
                if name in archive.namelist():
                    root = ET.fromstring(archive.read(name))
                    sections.append(name + ": " + " ".join(node.text or "" for node in root.iter() if node.tag.endswith("}t")))
        text = "\n".join(sections)
    else:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > 10:
                raise ValueError("Split PDFs into files of at most 10 pages.")
            pages = [p.extract_text(layout=True) or "" for p in pdf.pages]
            if any(len(p.strip()) < 20 for p in pages):
                return None  # OCR the full document, including mixed scanned/text pages.
            text = "\n".join(f"Page {i+1}\n{p}" for i, p in enumerate(pages))
    if not text.strip():
        raise ValueError("No readable text found in this document.")
    if len(text) > MAX_TEXT:
        raise ValueError("Document text exceeds the 100,000-character limit; split it to preserve every footnote.")
    return text
