"""Five real demo input files generated from the illustrative quote data."""
import io
import zipfile

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .sample import sample_workspace


def generate_quotation_template() -> bytes:
    """Editable supplier input with explicit units, terms and filling guidance."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    instructions = book.active
    instructions.title = "Instructions"
    instructions.append(["SUPPLIER QUOTATION TEMPLATE", "How to fill in your own data"])
    guidance = [
        ("1. Replace examples", "All supplier details, prices, quantities and certifications are examples. Replace them before uploading."),
        ("2. One file per supplier", "Make a separate copy for each supplier. Change Supplier name in Supplier Details."),
        ("3. Match item codes", "Use identical SKU / Item Code values for the same item across all suppliers. Row order may differ."),
        ("4. Fill Quotation", "Keep the header row. Replace/delete example rows and add as many item rows as needed."),
        ("5. Quantity", "Quantity (boxes) is a whole box count. Enter the requested quantity; leave unknown values blank."),
        ("6. Price unit", "Select per_box, per_kg, or per_1000_boxes. Unit price is the original quoted rate in that unit."),
        ("7. Weight for kg prices", "For per_kg quotes, provide buyer-approved kg/box weights in the app's Master SKUs. No weight is guessed."),
        ("8. Freight", "Use Freight per box for variable shipping and Fixed freight per order in Supplier Details for fixed charges. Enter 0 only if included/free; leave unknown charges blank."),
        ("9. Discounts", "Unconditional discount (%) is for an unconditional line discount only. Do not put conditional rebates in this column."),
        ("10. Conditional rebate", "Write every condition in Supplier Details: percentage, threshold, strictly above versus at least, whether all SKUs are required, and whether freight is discounted."),
        ("11. Currency and tax", "State the actual currency and GST percentage. Comparison defaults to GST excluded. Foreign currency requires review/conversion."),
        ("12. Certifications", "List actual supplier certifications. These do not automatically constitute buyer QA approval."),
        ("13. Upload", "Save as .xlsx, then select two or more supplier files together using Upload quotes. Review extracted values before confirming terms/QA."),
    ]
    for row in guidance:
        instructions.append(row)
    details = book.create_sheet("Supplier Details")
    details.append(["Field", "Value"])
    for row in [
        ("Supplier name", "Example Cartons — replace with your supplier"),
        ("Currency", "INR"), ("Payment terms", "Net 60 days"),
        ("Fixed freight per order (INR)", 0),
        ("GST percentage (extra, excluded from base rates)", 18),
        ("Certifications", "Example only: replace with actual certificates, or enter None"),
        ("Conditional rebate / other commercial conditions", "None. Replace with the full condition if a rebate exists."),
    ]:
        details.append(row)
    quotation = book.create_sheet("Quotation")
    quotation.append(["SKU / Item Code", "Description", "Quantity (boxes)", "Unit Price", "Currency", "Price Unit", "Freight per box (INR)", "Unconditional discount (%)", "Notes"])
    for row in [
        ("SKU-001", "Example 3-ply carton — replace description", 50000, 11.80, "INR", "per_box", 0, 0, "Example price: replace with actual quotation"),
        ("SKU-002", "Example 5-ply carton — replace description", 25000, 27.90, "INR", "per_box", 0, 0, "Example price: replace with actual quotation"),
        ("SKU-003", "Example 7-ply carton — replace description", 10000, 61.50, "INR", "per_box", 0, 0, "Example price: replace with actual quotation"),
    ]:
        quotation.append(row)
    unit_validation = DataValidation(type="list", formula1='"per_box,per_kg,per_1000_boxes"', allow_blank=True)
    unit_validation.errorTitle = "Choose a supported price unit"
    unit_validation.error = "Choose per_box, per_kg, or per_1000_boxes."
    unit_validation.showErrorMessage = True
    quotation.add_data_validation(unit_validation)
    unit_validation.add("F2:F5001")
    quotation.auto_filter.ref = quotation.dimensions
    for sheet in book:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="335F47")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[1].height = 32
    for sheet in (instructions, details):
        sheet.column_dimensions["A"].width = 48
        sheet.column_dimensions["B"].width = 110
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 44
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column, width in {"A": 22, "B": 48, "C": 22, "D": 16, "E": 14, "F": 23, "G": 24, "H": 26, "I": 54}.items():
        quotation.column_dimensions[column].width = width
    for row in quotation.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for index in (3, 6, 7):
            row[index].number_format = "0.00"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def generate_sample_zip() -> bytes:
    workspace = sample_workspace()
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        q = workspace.quotes[0]
        book = Workbook()
        sheet = book.active
        sheet.append([q.vendor_name, "Payment: Net 60. Freight included. GST excluded."])
        sheet.append(["Material ref", "Our offered rate", "Basis"])
        for line in q.line_items:
            sheet.append([line.sku_id, float(line.quoted_price), "INR per box"])
        data = io.BytesIO()
        book.save(data)
        archive.writestr(q.source_file, data.getvalue())
        q = workspace.quotes[1]
        data = io.BytesIO()
        pdf = canvas.Canvas(data, pagesize=A4)
        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(45, 795, q.vendor_name + " | Quotation")
        pdf.setFont("Helvetica", 11)
        pdf.drawString(45, 770, "Payment: Net 45. Currency: INR. GST excluded.")
        for i, line in enumerate(q.line_items):
            pdf.drawString(45, 745 - i * 19, f"{line.sku_id}    INR {line.quoted_price} per box")
        pdf.setFont("Helvetica", 8)
        pdf.drawString(45, 120, "Footnote: unconditional 5% discount on all listed rates. Freight INR 7,500 per awarded order.")
        pdf.save()
        archive.writestr(q.source_file, data.getvalue())
        q = workspace.quotes[2]
        doc = Document()
        doc.add_heading(q.vendor_name, 0)
        doc.add_paragraph("Net 60 payment. Freight included. GST excluded. Rates in INR.")
        for line in q.line_items:
            doc.add_paragraph(f"For your item {line.sku_id}, we offer a price of INR {line.quoted_price} for each box.")
        data = io.BytesIO()
        doc.save(data)
        archive.writestr(q.source_file, data.getvalue())
        q = workspace.quotes[3]
        image = Image.new("RGB", (1100, 1500), "#f4f0e7")
        draw = ImageDraw.Draw(image)
        try:
            font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 27)
        except OSError:
            font = ImageFont.load_default(size=27)
        draw.text((55, 45), q.vendor_name, font=font, fill="#222222")
        draw.text((55, 95), "Net 30. Freight included. GST excluded.", font=font, fill="#222222")
        for i, line in enumerate(q.line_items):
            draw.text((55, 155 + i * 40), f"{line.sku_id}       INR {line.quoted_price} / kg", font=font, fill="#222222")
        data = io.BytesIO()
        image.filter(ImageFilter.GaussianBlur(0.7)).save(data, format="PNG")
        archive.writestr(q.source_file, data.getvalue())
        archive.writestr(workspace.quotes[4].source_file, "From: Metro Cartons\nINR 42/kg for all 5-ply cartons, rest same as last year.\n")
        archive.writestr("README.txt", "Illustrative demo inputs, not actual supplier quotations. Load Sample Dataset uses curated extraction results and known demo QA decisions. Re-uploading these files invokes Sarvam; confirm QA and commercial terms manually. The Metro email intentionally has no historical rates.\n")
    return output.getvalue()
