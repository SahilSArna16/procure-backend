"""Synthetic INPUTS only. All production demo outputs use real Sarvam calls."""
import io
from decimal import Decimal
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

from .models import CurrencyConversion, QualityQuestion, RFx, VendorInvitation, Workspace
from .sample import sample_workspace
from . import settings

QUOTE_NAMES = ["Apex_30_Items.xlsx", "Packwell_Footnote.pdf", "GreenBox_27_Items_USD.docx", "Shree_Angled_Rate_Card.png", "Metro_Reply.txt"]
CERT_NAMES = {name: name.split('_')[0] + "_ISO_Evidence.pdf" for name in QUOTE_NAMES[:4]}
QUESTIONS = [
    QualityQuestion(id="iso9001", prompt="Is your ISO 9001 certification current? Attach certificate evidence.", expected_answer="yes", requires_attachment=True),
    QualityQuestion(id="fsc", prompt="Are the supplied cartons FSC compliant?", expected_answer="yes"),
    QualityQuestion(id="burst_factor", prompt="What burst factor (BF) do you guarantee?", answer_type="number", minimum="24"),
    QualityQuestion(id="lead_time", prompt="What is your lead time in days?", answer_type="number", maximum="10"),
]


def demo_workspace():
    master = sample_workspace().skus
    for sku in master:
        sku.pieces_per_box = Decimal("1")
    return Workspace(id=str(uuid4()), demo_mode=True, master_origin="buyer", skus=master, fx_rates=[CurrencyConversion(currency="USD", inr_per_unit="83", source="Synthetic demo assumption; not a live exchange rate", confirmed=False)])


def draft_seed(workspace):
    return RFx(title="Corrugated packaging RFQ · 30 factory SKUs", scope="Supply 30 carton specifications across Pune, Ahmedabad and Chennai. The buyer's approved master defines quantities, weights and one piece per box.", skus=workspace.skus, questionnaire=QUESTIONS, vendors=[VendorInvitation(name=name, email=name.split()[0].lower() + "@vendor.example") for name in ["Apex Containers", "Packwell Industries", "GreenBox Packaging", "Shree Corrugators", "Metro Cartons"]], assumptions=["Synthetic buyer master, quantities and material weights; review before use.", "USD conversion is indicative and requires explicit buyer confirmation."])


def make_files():
    master = demo_workspace().skus
    rates = [[(Decimal("12") + Decimal(i) * Decimal(".87") + Decimal((i * 7 + v * 11) % 9) / 4).quantize(Decimal(".01")) for i in range(30)] for v in range(4)]
    outputs = {}
    book = Workbook()
    sheet = book.active
    sheet.title = "Supplier Response"
    sheet.append(["Vendor: Apex Containers", "Net 60 days. Freight included. Currency: INR. GST 18% extra. No discounts."])
    sheet.append(["Item Code", "Description", "Order Batch Size", "Unit Price (INR/Box)"])
    for i in list(range(10, 20)) + list(range(0, 10)) + list(range(20, 30)):
        sku = master[i]
        sheet.append([sku.sku_id, sku.description, sku.quantity, float(rates[0][i])])
    sheet.append(["Quality responses", "ISO 9001: YES; FSC: YES; burst factor: 24; lead time: 7 days. ISO evidence attached separately."])
    out = io.BytesIO(); book.save(out); outputs[QUOTE_NAMES[0]] = out.getvalue()
    out = io.BytesIO(); pdf = canvas.Canvas(out, pagesize=A4)
    for page in range(2):
        pdf.setFont("Helvetica-Bold", 16); pdf.drawString(40, 795, "Packwell Industries | Commercial Quotation")
        pdf.setFont("Helvetica", 10); pdf.drawString(40, 770, "Currency: INR. Net 45 days. GST 18% extra. Freight INR 7,500 per order.")
        pdf.drawString(40, 755, "No per-box freight charge. Fixed freight only. No unconditional line discounts.")
        pdf.drawString(40, 738, "SKU ID     Quantity      INR / box")
        for j in range(15):
            i = page * 15 + j
            pdf.drawString(40, 725 - j * 26, f"{master[i].sku_id}    {master[i].quantity} boxes    INR {rates[1][i]} per box")
        pdf.setFont("Helvetica", 8)
        pdf.drawString(40, 245, "Footnote: 5% goods rebate only if base goods value exceeds INR 6,000,000 and all 30 SKUs")
        pdf.drawString(40, 231, "are awarded in a single order. No rebate for partial awards or exact threshold. Freight not discounted.")
        pdf.drawString(40, 203, "ISO 9001: YES. FSC: YES. Guaranteed burst factor: 26. Lead time: 8 days.")
        pdf.drawString(40, 189, "ISO evidence attached separately. SAMPLE DATA, NOT A REAL OFFER.")
        pdf.showPage()
    pdf.save(); outputs[QUOTE_NAMES[1]] = out.getvalue()
    doc = Document(); doc.add_heading("GreenBox Packaging", 0)
    doc.add_paragraph("Currency: USD. Prices per 100 pieces. One piece means one completed carton. Net 60 days. Freight included. GST 18% extra. No discounts.")
    for i in range(27):
        amount = (rates[2][i] / 83 * 100).quantize(Decimal(".01"))
        doc.add_paragraph(f"For {master[i].sku_id}, we offer USD {amount} per 100 pieces for {master[i].quantity} boxes. Description: {master[i].description}.")
    doc.add_paragraph("SKU-028, SKU-029 and SKU-030: no quotation this round. Quality: ISO 9001 YES; FSC NO; burst factor 20; lead time 12 days. ISO certificate attached. Synthetic sample only.")
    out = io.BytesIO(); doc.save(out); outputs[QUOTE_NAMES[2]] = out.getvalue()
    image = Image.new("RGB", (1300, 1770), "#f5f0e3"); draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(settings.ROOT / "assets/fonts/DejaVuSans.ttf"), 26)
    draw.text((45, 40), "SHREE CORRUGATORS | RATE CARD", font=font, fill="#222")
    draw.text((45, 88), "INR per kg. Net 30 days. Freight included. GST 18% extra.", font=font, fill="#222")
    for i, sku in enumerate(master):
        rate = (rates[3][i] / sku.weight_kg).quantize(Decimal(".01"))
        draw.text((45, 148 + i * 44), f"{sku.sku_id}     INR {rate} / kg     {sku.quantity} boxes", font=font, fill="#222")
    draw.text((45, 1490), "ISO 9001 YES. FSC YES. Burst factor 24. Lead time 9 days.", font=font, fill="#222")
    draw.text((45, 1540), "No discounts. ISO evidence attached. Synthetic sample.", font=font, fill="#222")
    out = io.BytesIO(); image.filter(ImageFilter.GaussianBlur(.35)).rotate(7, expand=True, fillcolor="#d6d2c9").save(out, "PNG"); outputs[QUOTE_NAMES[3]] = out.getvalue()
    outputs[QUOTE_NAMES[4]] = b"From: Metro Cartons\nINR 42/kg for all 5-ply, INR 38/kg for all 3-ply, rest same as last year. Freight extra, amount to follow. Net 60 days. GST extra.\nISO 9001 yes; FSC and burst factor confirmation pending; lead time to follow. No certificate attached.\nSynthetic sample email.\n"
    for name, cert_name in CERT_NAMES.items():
        out = io.BytesIO(); cert = canvas.Canvas(out, pagesize=A4)
        cert.setFont("Helvetica-Bold", 20); cert.drawString(45, 780, "SYNTHETIC CERTIFICATE EVIDENCE")
        cert.setFont("Helvetica", 12)
        for y, text in [(730, "Supplier: " + {QUOTE_NAMES[0]: "Apex Containers", QUOTE_NAMES[1]: "Packwell Industries", QUOTE_NAMES[2]: "GreenBox Packaging", QUOTE_NAMES[3]: "Shree Corrugators"}[name]), (695, "Standard: ISO 9001:2015"), (660, "Demo validity: 01 January 2026 to 31 December 2027"), (625, "Issuer: Example Quality Registrar (fictional)"), (590, "This is demonstration evidence, not a genuine certificate.")]:
            cert.drawString(45, y, text)
        cert.save(); outputs[cert_name] = out.getvalue()
    return outputs


def demo_zip():
    out = io.BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        for name, data in make_files().items():
            archive.writestr(name, data)
        archive.writestr("README.txt", "Synthetic input dataset only. Use real AI extraction for all five quotations. Upload ISO evidence against the iso9001 question, review it, approve the buyer quality rules, and confirm commercial terms. The USD rate is indicative and unconfirmed; the Word response omits three items, and Metro has unknown freight/historical prices.\n")
    return out.getvalue()
